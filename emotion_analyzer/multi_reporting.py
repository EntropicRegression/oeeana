from __future__ import annotations

import re
from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import RunAnalysis, SegmentObservation, make_run_analysis
from .reporting import (
    _save_atomic,
    _write_info_sheet,
    _write_raw_sheet,
    _write_summary_sheet,
    read_run_report,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, color="FFFFFF", bold=True)
BODY_FONT = Font(name="Arial", size=10, color="1F1F1F")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F1F1F")


class MultiSetReportError(ValueError):
    """Raised when per-set reports cannot be combined safely."""


def _style_header(worksheet, row: int, first_column: int, last_column: int) -> None:
    for column in range(first_column, last_column + 1):
        cell = worksheet.cell(row, column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fit_columns(worksheet, maximum: int = 36) -> None:
    for column_cells in worksheet.columns:
        for cell in column_cells:
            if cell.value is not None:
                cell.font = cell.font.copy(name="Arial") if cell.font else BODY_FONT
        letter = get_column_letter(column_cells[0].column)
        width = max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2)
        worksheet.column_dimensions[letter].width = min(maximum, width)


def _safe_sheet_name(value: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\\/*?:\[\]]+", "_", value).strip() or "Set"
    base = cleaned[:31]
    candidate = base
    suffix = 2
    while candidate.casefold() in used:
        tail = f"_{suffix}"
        candidate = base[: 31 - len(tail)] + tail
        suffix += 1
    used.add(candidate.casefold())
    return candidate


def _subject_map(run: RunAnalysis, set_name: str):
    subjects = {}
    duplicates: list[str] = []
    for subject in run.comparison_subjects:
        if subject.subject_id in subjects:
            duplicates.append(subject.subject_id)
        subjects[subject.subject_id] = subject
    if duplicates:
        duplicate_text = ", ".join(sorted(set(duplicates)))
        raise MultiSetReportError(
            f"Analysis Set '{set_name}' 有重複受試者 ID：{duplicate_text}"
        )
    return subjects


def build_multi_set_analysis(
    set_reports: Sequence[tuple[str, Path]],
) -> tuple[RunAnalysis, tuple[tuple[str, RunAnalysis], ...]]:
    """Represent each Analysis Set mean as one high-level segment per subject."""
    if not set_reports:
        raise MultiSetReportError("至少需要一份 Analysis Set 報表。")

    names = [name.strip() for name, _ in set_reports]
    if any(not name for name in names):
        raise MultiSetReportError("Analysis Set 名稱不可為空白。")
    if len({name.casefold() for name in names}) != len(names):
        raise MultiSetReportError("Analysis Set 名稱不可重複。")

    loaded: list[tuple[str, RunAnalysis]] = []
    maps: list[dict[str, object]] = []
    subject_order: list[str] = []
    seen_subjects: set[str] = set()

    for set_name, report_path in set_reports:
        run = read_run_report(Path(report_path))
        loaded.append((set_name, run))
        subject_map = _subject_map(run, set_name)
        maps.append(subject_map)
        for subject in run.comparison_subjects:
            if subject.subject_id not in seen_subjects:
                seen_subjects.add(subject.subject_id)
                subject_order.append(subject.subject_id)

    observations: list[SegmentObservation] = []
    ordered_subjects: list[tuple[str, str, bool, str | None]] = []

    for subject_id in subject_order:
        first_subject = next(
            (
                subject_map[subject_id]
                for subject_map in maps
                if subject_id in subject_map
            ),
            None,
        )
        audio_name = (
            getattr(first_subject, "audio_name", None)
            or f"{subject_id}-combined"
        )
        ordered_subjects.append((subject_id, str(audio_name), False, None))

        for set_index, ((set_name, _run), subject_map) in enumerate(
            zip(loaded, maps), 1
        ):
            subject = subject_map.get(subject_id)
            if subject is None:
                observations.append(
                    SegmentObservation(
                        subject_id=subject_id,
                        audio_name=str(audio_name),
                        is_reference=False,
                        segment_index=set_index,
                        error=f"缺少 Analysis Set：{set_name}",
                    )
                )
                continue

            mean_l1 = getattr(subject, "mean_l1", None)
            mean_l2 = getattr(subject, "mean_l2", None)
            error = getattr(subject, "error", None)
            if mean_l1 is None and mean_l2 is None and not error:
                error = f"Analysis Set {set_name} 沒有有效距離數值"
            observations.append(
                SegmentObservation(
                    subject_id=subject_id,
                    audio_name=str(audio_name),
                    is_reference=False,
                    segment_index=set_index,
                    l1_distance=mean_l1,
                    l2_distance=mean_l2,
                    error=error,
                )
            )

    analysis = make_run_analysis(
        observations,
        "MULTI_SET",
        "oeeana-multi-set",
        subject_order=ordered_subjects,
    )
    return analysis, tuple(loaded)


def _write_overview(
    workbook: Workbook,
    analysis: RunAnalysis,
    set_names: Sequence[str],
) -> None:
    worksheet = workbook.active
    worksheet.title = "整體摘要"
    worksheet.sheet_view.showGridLines = False
    worksheet["A2"] = "多段 Analysis Set 整體摘要"
    worksheet["A2"].font = TITLE_FONT
    worksheet["A3"] = "Overall 為各 Analysis Set 平均值的等權平均。"
    worksheet["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")

    headers = ["受試者", "有效 Set 數"]
    for set_name in set_names:
        headers.extend([f"{set_name} L1", f"{set_name} L2"])
    headers.extend(["Overall L1", "Overall L2", "缺失/錯誤 Set"])
    header_row = 5
    for column, value in enumerate(headers, 1):
        worksheet.cell(header_row, column, value)
    _style_header(worksheet, header_row, 1, len(headers))

    for subject in analysis.comparison_subjects:
        by_index = {segment.segment_index: segment for segment in subject.segments}
        values: list[object] = [subject.subject_id, subject.valid_segments]
        issues: list[str] = []
        for index, set_name in enumerate(set_names, 1):
            segment = by_index.get(index)
            values.extend(
                [
                    segment.l1_distance if segment else None,
                    segment.l2_distance if segment else None,
                ]
            )
            if segment is None or segment.error:
                issues.append(set_name)
        values.extend(
            [
                subject.mean_l1,
                subject.mean_l2,
                "、".join(issues) if issues else None,
            ]
        )
        worksheet.append(values)

    worksheet.freeze_panes = "A6"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row[2:-1]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"
    _fit_columns(worksheet, maximum=28)


def _write_set_sheet(workbook: Workbook, set_name: str, run: RunAnalysis, used: set[str]) -> None:
    worksheet = workbook.create_sheet(_safe_sheet_name(set_name, used))
    worksheet.sheet_view.showGridLines = False
    worksheet["A2"] = f"Analysis Set：{set_name}"
    worksheet["A2"].font = TITLE_FONT
    headers = [
        "受試者",
        "音檔名稱",
        "有效段數",
        "L1平均截距",
        "L1總計截距",
        "L2平均差距",
        "L2總計差距",
        "錯誤",
    ]
    header_row = 4
    for column, value in enumerate(headers, 1):
        worksheet.cell(header_row, column, value)
    _style_header(worksheet, header_row, 1, len(headers))
    for subject in run.comparison_subjects:
        worksheet.append(
            [
                subject.subject_id,
                subject.audio_name,
                subject.valid_segments,
                subject.mean_l1,
                subject.total_l1,
                subject.mean_l2,
                subject.total_l2,
                subject.error,
            ]
        )
    worksheet.freeze_panes = "A5"
    worksheet.auto_filter.ref = worksheet.dimensions
    for row in worksheet.iter_rows(min_row=header_row + 1):
        for cell in row[3:7]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"
    _fit_columns(worksheet)


def _write_set_index(workbook: Workbook, loaded: Sequence[tuple[str, RunAnalysis]]) -> None:
    worksheet = workbook.create_sheet("Set對照")
    worksheet.append(["高階段落", "Analysis Set", "原始段落數", "標準音檔"])
    _style_header(worksheet, 1, 1, 4)
    for index, (set_name, run) in enumerate(loaded, 1):
        worksheet.append([index, set_name, run.segment_count, run.reference_audio])
    worksheet.freeze_panes = "A2"
    _fit_columns(worksheet, maximum=48)


def write_multi_set_report(
    path: Path,
    set_reports: Sequence[tuple[str, Path]],
) -> RunAnalysis:
    """Combine per-set reports into one subject-level report readable by the comparison UI."""
    analysis, loaded = build_multi_set_analysis(set_reports)
    workbook = Workbook()
    set_names = [name for name, _ in loaded]
    _write_overview(workbook, analysis, set_names)

    used = {"整體摘要".casefold()}
    for set_name, run in loaded:
        _write_set_sheet(workbook, set_name, run, used)
    _write_set_index(workbook, loaded)

    # These standard sheets keep the combined workbook compatible with read_run_report()
    # and therefore with the existing pre/post comparison interface.
    _write_summary_sheet(workbook, analysis)
    _write_raw_sheet(workbook, analysis)
    _write_info_sheet(
        workbook,
        analysis,
        {
            "report_type": "multi_set_run_analysis",
            "analysis_set_count": len(loaded),
            "analysis_sets": " | ".join(set_names),
            "overall_aggregation": "equal_weight_mean_of_set_means",
        },
    )
    _save_atomic(workbook, Path(path))
    return analysis
