from __future__ import annotations

import math
import os
import re
from datetime import datetime
from pathlib import Path
from typing import Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .analytics import (
    EMOTION_LABELS,
    SCHEMA_VERSION,
    ComparisonStatistics,
    ExperimentComparison,
    GroupEffectTestResult,
    RunAnalysis,
    SegmentObservation,
    make_run_analysis,
    normalize_subject_id,
    summarize_run,
)


HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(name="Arial", size=10, color="FFFFFF", bold=True)
STANDARD_FILL = PatternFill("solid", fgColor="EAF2F8")
ERROR_FILL = PatternFill("solid", fgColor="FCE8E6")
TITLE_FONT = Font(name="Arial", size=14, bold=True, color="1F1F1F")
BODY_FONT = Font(name="Arial", size=10, color="1F1F1F")


class ReportFormatError(ValueError):
    """Raised when an Excel file is not a supported analysis report."""


def _number(value: object) -> float | None:
    try:
        result = float(value)
    except (TypeError, ValueError):
        return None
    return result if math.isfinite(result) else None


def _style_header(worksheet, row: int, first_column: int, last_column: int) -> None:
    for column in range(first_column, last_column + 1):
        cell = worksheet.cell(row, column)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _fit_columns(worksheet, *, maximum: int = 42) -> None:
    for column_cells in worksheet.columns:
        for cell in column_cells:
            if cell.value is not None and cell.font.name != "Arial":
                cell.font = BODY_FONT
                cell.alignment = Alignment(vertical="center", wrap_text=False)
        column_letter = get_column_letter(column_cells[0].column)
        width = max(10, max(len(str(cell.value or "")) for cell in column_cells) + 2)
        worksheet.column_dimensions[column_letter].width = min(maximum, width)


def _save_atomic(workbook: Workbook, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(path.suffix + ".tmp")
    try:
        workbook.save(temporary)
        os.replace(temporary, path)
    except Exception:
        try:
            temporary.unlink(missing_ok=True)
        except OSError:
            pass
        raise


def _write_legacy_sheet(workbook: Workbook, rows: Sequence[object], segment_count: int) -> None:
    worksheet = workbook.active
    worksheet.title = "情緒分析結果"
    worksheet.sheet_view.showGridLines = False
    headers = ["音檔名稱"]
    for index in range(1, segment_count + 1):
        headers.extend(
            [
                f"第{index}段情緒原始分數",
                f"第{index}段主情緒",
                f"第{index}段跟標準音檔分析的差距",
            ]
        )
    worksheet.append(headers)
    for row_index, row in enumerate(rows):
        values: list[object] = [getattr(row, "audio_name", "")]
        segments = getattr(row, "segments", ())
        row_error = getattr(row, "error", None)
        for index in range(segment_count):
            segment = segments[index] if index < len(segments) else None
            segment_error = getattr(segment, "error", None) if segment is not None else None
            if row_error or segment is None or segment_error:
                values.extend([None, row_error or segment_error or "分析失敗", None])
                continue
            distance = 0.0 if row_index == 0 else _number(getattr(segment, "distance", None))
            values.extend(
                [
                    _number(getattr(segment, "top_score", None)),
                    getattr(segment, "emotion", None),
                    distance,
                ]
            )
        worksheet.append(values)

    _style_header(worksheet, 1, 1, len(headers))
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    for cell in worksheet[2]:
        cell.fill = STANDARD_FILL
    for row in worksheet.iter_rows(min_row=2, min_col=2):
        for cell in row:
            if (cell.column - 2) % 3 in (0, 2) and isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"
    _fit_columns(worksheet)


def _write_summary_sheet(workbook: Workbook, analysis: RunAnalysis) -> None:
    worksheet = workbook.create_sheet("分析摘要")
    worksheet.sheet_view.showGridLines = False
    worksheet["A2"] = "音檔情緒分析摘要"
    worksheet["A2"].font = TITLE_FONT
    worksheet["A3"] = "L1 為 Friendly Support 相容截距；L2 為原有 oeeana 差距。"
    worksheet["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")
    headers = [
        "受試者",
        "音檔名稱",
        "標準音檔",
        "有效段數",
        "L1平均截距",
        "L1總計截距",
        "L2平均差距",
        "L2總計差距",
        "L1最大偏離段落",
        "L2最大偏離段落",
        "錯誤",
    ]
    header_row = 5
    for column, value in enumerate(headers, 1):
        worksheet.cell(header_row, column, value)
    for subject in analysis.subjects:
        worksheet.append(
            [
                subject.subject_id,
                subject.audio_name,
                "是" if subject.is_reference else "否",
                subject.valid_segments,
                subject.mean_l1,
                subject.total_l1,
                subject.mean_l2,
                subject.total_l2,
                subject.max_l1_segment,
                subject.max_l2_segment,
                subject.error,
            ]
        )
    last_summary_row = header_row + len(analysis.subjects)
    _style_header(worksheet, header_row, 1, len(headers))
    worksheet.auto_filter.ref = f"A{header_row}:K{last_summary_row}"
    worksheet.freeze_panes = "A6"
    for row in worksheet.iter_rows(min_row=header_row + 1, max_row=last_summary_row):
        if row[2].value == "是":
            for cell in row:
                cell.fill = STANDARD_FILL
        if row[10].value:
            row[10].fill = ERROR_FILL
        for cell in row[4:8]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"

    comparison_subjects = analysis.comparison_subjects
    trend_header_row = last_summary_row + 3
    if comparison_subjects and analysis.segment_count:
        worksheet.cell(trend_header_row - 1, 1, "各段 L1 情緒截距")
        worksheet.cell(trend_header_row - 1, 1).font = Font(name="Arial", size=11, bold=True)
        trend_headers = ["段落"] + [subject.subject_id for subject in comparison_subjects]
        for column, value in enumerate(trend_headers, 1):
            worksheet.cell(trend_header_row, column, value)
        _style_header(worksheet, trend_header_row, 1, len(trend_headers))
        for segment_index in range(1, analysis.segment_count + 1):
            values: list[object] = [segment_index]
            for subject in comparison_subjects:
                observation = next(
                    (item for item in subject.segments if item.segment_index == segment_index),
                    None,
                )
                values.append(observation.l1_distance if observation else None)
            worksheet.append(values)
        trend_last_row = trend_header_row + analysis.segment_count
        for row in worksheet.iter_rows(
            min_row=trend_header_row + 1,
            max_row=trend_last_row,
            min_col=2,
            max_col=len(trend_headers),
        ):
            for cell in row:
                if isinstance(cell.value, (int, float)):
                    cell.number_format = "0.000000"
        chart = LineChart()
        chart.title = "各段 L1 情緒截距"
        chart.y_axis.title = "L1 截距"
        chart.x_axis.title = "段落"
        chart.style = 13
        chart.height = 8
        chart.width = 15
        chart.add_data(
            Reference(
                worksheet,
                min_col=2,
                max_col=len(trend_headers),
                min_row=trend_header_row,
                max_row=trend_last_row,
            ),
            titles_from_data=True,
        )
        chart.set_categories(
            Reference(worksheet, min_col=1, min_row=trend_header_row + 1, max_row=trend_last_row)
        )
        worksheet.add_chart(chart, "M2")

    _fit_columns(worksheet)


def _write_raw_sheet(workbook: Workbook, analysis: RunAnalysis) -> None:
    worksheet = workbook.create_sheet("原始分數")
    worksheet.sheet_view.showGridLines = False
    headers = [
        "受試者",
        "音檔名稱",
        "標準音檔",
        "段落",
        "開始秒",
        "結束秒",
        "匹配信心",
        "主情緒",
        "主情緒分數",
        "前三高",
        "L1情緒截距",
        "L2情緒差距",
        "最大變化情緒",
        "變化方向",
        "變化量",
        *[f"分數_{label}" for label in EMOTION_LABELS],
        "錯誤",
    ]
    worksheet.append(headers)
    for subject in analysis.subjects:
        for segment in subject.segments:
            scores = list(segment.scores) if segment.scores is not None else [None] * len(EMOTION_LABELS)
            top_three = ", ".join(f"{label}:{score:.4f}" for label, score in segment.top_three)
            worksheet.append(
                [
                    subject.subject_id,
                    subject.audio_name,
                    "是" if subject.is_reference else "否",
                    segment.segment_index,
                    segment.start,
                    segment.end,
                    segment.confidence,
                    segment.main_emotion,
                    segment.top_score,
                    top_three,
                    segment.l1_distance,
                    segment.l2_distance,
                    segment.max_change_emotion,
                    segment.max_change_direction,
                    segment.max_change_amount,
                    *scores,
                    segment.error,
                ]
            )

    _style_header(worksheet, 1, 1, len(headers))
    worksheet.freeze_panes = "D2"
    worksheet.auto_filter.ref = worksheet.dimensions
    numeric_headers = {
        "開始秒",
        "結束秒",
        "匹配信心",
        "主情緒分數",
        "L1情緒截距",
        "L2情緒差距",
        "變化量",
        *[f"分數_{label}" for label in EMOTION_LABELS],
    }
    for column, header in enumerate(headers, 1):
        if header in numeric_headers:
            for cell in worksheet.iter_cols(
                min_col=column,
                max_col=column,
                min_row=2,
                max_row=worksheet.max_row,
            ):
                for item in cell:
                    if isinstance(item.value, (int, float)):
                        item.number_format = "0.000000"
    _fit_columns(worksheet, maximum=32)


def _write_info_sheet(
    workbook: Workbook,
    analysis: RunAnalysis,
    metadata: Mapping[str, object] | None,
) -> None:
    worksheet = workbook.create_sheet("報表資訊")
    worksheet.sheet_view.showGridLines = False
    rows = [
        ("欄位", "值"),
        ("schema_version", SCHEMA_VERSION),
        ("report_type", "run_analysis"),
        ("generated_at", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("reference_audio", analysis.reference_audio),
        ("segment_count", analysis.segment_count),
        ("legacy_distance_metric", "l2"),
        ("friendly_distance_metric", "l1"),
        ("source_format", analysis.source_format),
    ]
    for key, value in (metadata or {}).items():
        rows.append((str(key), value))
    for row in rows:
        worksheet.append(row)
    _style_header(worksheet, 1, 1, 2)
    _fit_columns(worksheet, maximum=52)
    worksheet.column_dimensions["A"].width = 28
    worksheet.column_dimensions["B"].width = 48


def write_run_report(
    path: Path,
    rows: Sequence[object],
    segment_count: int,
    metadata: Mapping[str, object] | None = None,
) -> RunAnalysis:
    """Write the preserved legacy layout plus summary, raw data and metadata sheets."""
    analysis = summarize_run(rows)
    workbook = Workbook()
    _write_legacy_sheet(workbook, rows, segment_count)
    _write_summary_sheet(workbook, analysis)
    _write_raw_sheet(workbook, analysis)
    _write_info_sheet(workbook, analysis, metadata)
    _save_atomic(workbook, path)
    return analysis


def _sheet_records(worksheet, header_row: int = 1) -> list[dict[str, object]]:
    headers = [cell.value for cell in worksheet[header_row]]
    records: list[dict[str, object]] = []
    for values in worksheet.iter_rows(min_row=header_row + 1, values_only=True):
        if not any(value is not None for value in values):
            continue
        records.append({str(header): value for header, value in zip(headers, values) if header is not None})
    return records


def _information(workbook) -> dict[str, object]:
    if "報表資訊" not in workbook.sheetnames:
        return {}
    return {
        str(key): value
        for key, value in workbook["報表資訊"].iter_rows(min_row=2, max_col=2, values_only=True)
        if key is not None
    }


def _read_v2(workbook) -> RunAnalysis:
    info = _information(workbook)
    records = _sheet_records(workbook["原始分數"])
    observations: list[SegmentObservation] = []
    for record in records:
        scores = tuple(_number(record.get(f"分數_{label}")) for label in EMOTION_LABELS)
        score_values = tuple(value for value in scores if value is not None)
        numeric_scores = tuple(float(value) for value in scores) if len(score_values) == len(EMOTION_LABELS) else None
        top_three = (
            tuple(sorted(zip(EMOTION_LABELS, numeric_scores), key=lambda pair: pair[1], reverse=True)[:3])
            if numeric_scores
            else ()
        )
        observations.append(
            SegmentObservation(
                subject_id=normalize_subject_id(
                    str(record.get("受試者") or record.get("音檔名稱") or "")
                ),
                audio_name=str(record.get("音檔名稱") or ""),
                is_reference=str(record.get("標準音檔") or "") == "是",
                segment_index=int(record.get("段落") or 0),
                start=_number(record.get("開始秒")),
                end=_number(record.get("結束秒")),
                confidence=_number(record.get("匹配信心")),
                main_emotion=str(record.get("主情緒") or "") or None,
                top_score=_number(record.get("主情緒分數")),
                scores=numeric_scores,
                top_three=top_three,
                l1_distance=_number(record.get("L1情緒截距")),
                l2_distance=_number(record.get("L2情緒差距")),
                max_change_emotion=str(record.get("最大變化情緒") or "") or None,
                max_change_direction=str(record.get("變化方向") or "") or None,
                max_change_amount=_number(record.get("變化量")),
                error=str(record.get("錯誤") or "") or None,
            )
        )

    order: list[tuple[str, str, bool, str | None]] = []
    seen: set[str] = set()
    for observation in observations:
        if observation.subject_id not in seen:
            seen.add(observation.subject_id)
            order.append(
                (
                    observation.subject_id,
                    observation.audio_name,
                    observation.is_reference,
                    None,
                )
            )
    reference_audio = str(info.get("reference_audio") or next(
        (item.audio_name for item in observations if item.is_reference),
        "標準音檔",
    ))
    return make_run_analysis(
        observations,
        reference_audio,
        str(info.get("source_format") or "oeeana-v2"),
        schema_version=str(info.get("schema_version") or SCHEMA_VERSION),
        subject_order=order,
    )


def _read_friendly_legacy(worksheet) -> RunAnalysis:
    records = _sheet_records(worksheet)
    headers = [str(cell.value or "") for cell in worksheet[1]]
    segment_columns: dict[int, str] = {}
    for header in headers:
        match = re.fullmatch(r"段落(\d+)_截距", header)
        if match:
            segment_columns[int(match.group(1))] = header
    observations: list[SegmentObservation] = []
    order: list[tuple[str, str, bool, str | None]] = []
    for record in records:
        name = str(record.get("人名") or "")
        is_reference = name == "[標準基準值]"
        subject_id = name if is_reference else normalize_subject_id(name)
        order.append((subject_id, name, is_reference, None))
        for index, column in sorted(segment_columns.items()):
            observations.append(
                SegmentObservation(
                    subject_id,
                    name,
                    is_reference,
                    index,
                    l1_distance=_number(record.get(column)),
                )
            )
    reference = next((name for name, _, is_reference, _ in order if is_reference), "[標準基準值]")
    return make_run_analysis(
        observations,
        reference,
        "friendly-support-legacy",
        schema_version="1",
        subject_order=order,
    )


def _read_oeeana_legacy(worksheet) -> RunAnalysis:
    records = _sheet_records(worksheet)
    headers = [str(cell.value or "") for cell in worksheet[1]]
    segment_indices = sorted(
        {
            int(match.group(1))
            for header in headers
            if (match := re.fullmatch(r"第(\d+)段跟標準音檔分析的差距", header))
        }
    )
    observations: list[SegmentObservation] = []
    order: list[tuple[str, str, bool, str | None]] = []
    for row_index, record in enumerate(records):
        audio_name = str(record.get("音檔名稱") or "")
        subject_id = normalize_subject_id(audio_name)
        is_reference = row_index == 0
        order.append((subject_id, audio_name, is_reference, None))
        for index in segment_indices:
            observations.append(
                SegmentObservation(
                    subject_id,
                    audio_name,
                    is_reference,
                    index,
                    main_emotion=str(record.get(f"第{index}段主情緒") or "") or None,
                    top_score=_number(record.get(f"第{index}段情緒原始分數")),
                    l2_distance=_number(record.get(f"第{index}段跟標準音檔分析的差距")),
                )
            )
    reference = order[0][1] if order else "標準音檔"
    return make_run_analysis(
        observations,
        reference,
        "oeeana-legacy",
        schema_version="1",
        subject_order=order,
    )


def read_run_report(path: Path) -> RunAnalysis:
    """Read current, oeeana legacy or Friendly Support legacy analysis workbooks."""
    if not path.is_file():
        raise ReportFormatError(f"找不到 Excel 報表：{path}")
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:
        raise ReportFormatError(f"無法讀取 Excel 報表：{path.name}") from exc
    if "原始分數" in workbook.sheetnames:
        return _read_v2(workbook)
    worksheet = workbook[workbook.sheetnames[0]]
    headers = {str(cell.value or "") for cell in worksheet[1]}
    if "人名" in headers and "平均差值" in headers:
        return _read_friendly_legacy(worksheet)
    if "音檔名稱" in headers and any("跟標準音檔分析的差距" in header for header in headers):
        return _read_oeeana_legacy(worksheet)
    raise ReportFormatError("Excel 不是可識別的 oeeana 或 Friendly Support 分析報表。")


def _statistics_row(statistics: ComparisonStatistics) -> list[object]:
    return [
        statistics.name,
        statistics.first.count,
        statistics.first.mean,
        statistics.first.standard_deviation,
        statistics.second.count,
        statistics.second.mean,
        statistics.second.standard_deviation,
        statistics.paired_test.paired_count,
        statistics.paired_test.t_statistic,
        statistics.paired_test.p_value,
        statistics.paired_test.note,
    ]


def _group_effect_row(result: GroupEffectTestResult) -> list[object]:
    return [
        result.name,
        result.experimental_count,
        result.experimental_mean_change,
        result.control_count,
        result.control_mean_change,
        result.effect_definition,
        result.effect,
        result.standard_error,
        result.t_statistic,
        result.degrees_freedom,
        result.p_value,
        result.confidence_interval_low,
        result.confidence_interval_high,
        result.note,
    ]


def write_comparison_report(path: Path, comparison: ExperimentComparison) -> None:
    workbook = Workbook()
    summary = workbook.active
    summary.title = "比較摘要"
    summary.sheet_view.showGridLines = False
    summary["A2"] = "前後測情緒比較"
    summary["A2"].font = TITLE_FONT
    summary["A3"] = f"比較指標：{comparison.metric.display_name}"
    summary["A3"].font = Font(name="Arial", size=10, italic=True, color="666666")
    headers = [
        "範圍",
        "前測樣本數",
        "前測平均",
        "前測標準差",
        "後測樣本數",
        "後測平均",
        "後測標準差",
        "成對樣本數",
        "t值",
        "p值",
        "結論",
    ]
    header_row = 5
    for column, value in enumerate(headers, 1):
        summary.cell(header_row, column, value)
    statistics_rows = [comparison.overall, *comparison.groups]
    for statistics in statistics_rows:
        summary.append(_statistics_row(statistics))
    _style_header(summary, header_row, 1, len(headers))
    for row in summary.iter_rows(min_row=header_row + 1, max_row=header_row + len(statistics_rows)):
        for cell in row[2:10]:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"

    group_test_title_row = header_row + len(statistics_rows) + 3
    summary.cell(group_test_title_row, 1, "組間效果檢定")
    summary.cell(group_test_title_row, 1).font = Font(name="Arial", size=12, bold=True, color="1F1F1F")
    group_test_header_row = group_test_title_row + 1
    group_test_headers = [
        "檢定",
        "實驗組n",
        "實驗組平均變化",
        "對照組n",
        "對照組平均變化",
        "效果定義",
        "效果估計",
        "標準誤",
        "t值",
        "自由度",
        "p值",
        "95%CI下限",
        "95%CI上限",
        "結論",
    ]
    for column, value in enumerate(group_test_headers, 1):
        summary.cell(group_test_header_row, column, value)
    for result in comparison.group_effect_tests:
        summary.append(_group_effect_row(result))
    _style_header(summary, group_test_header_row, 1, len(group_test_headers))
    for row in summary.iter_rows(
        min_row=group_test_header_row + 1,
        max_row=group_test_header_row + len(comparison.group_effect_tests),
    ):
        for column in (3, 5, 7, 8, 9, 10, 11, 12, 13):
            cell = row[column - 1]
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"

    chart_data_row = group_test_header_row + len(comparison.group_effect_tests) + 3
    summary.cell(chart_data_row, 1, "範圍")
    summary.cell(chart_data_row, 2, comparison.first_source)
    summary.cell(chart_data_row, 3, comparison.second_source)
    for offset, statistics in enumerate(statistics_rows, 1):
        summary.cell(chart_data_row + offset, 1, statistics.name)
        summary.cell(chart_data_row + offset, 2, statistics.first.mean)
        summary.cell(chart_data_row + offset, 3, statistics.second.mean)
    _style_header(summary, chart_data_row, 1, 3)
    chart = BarChart()
    chart.type = "col"
    chart.style = 10
    chart.title = f"{comparison.metric.display_name}平均比較"
    chart.y_axis.title = comparison.metric.display_name
    chart.x_axis.title = "範圍"
    chart.height = 8
    chart.width = 14
    chart.add_data(
        Reference(
            summary,
            min_col=2,
            max_col=3,
            min_row=chart_data_row,
            max_row=chart_data_row + len(statistics_rows),
        ),
        titles_from_data=True,
    )
    chart.set_categories(
        Reference(
            summary,
            min_col=1,
            min_row=chart_data_row + 1,
            max_row=chart_data_row + len(statistics_rows),
        )
    )
    summary.add_chart(chart, "P2")
    _fit_columns(summary)
    summary.column_dimensions["F"].width = 34
    summary.column_dimensions["N"].width = 42
    for row_number in range(
        group_test_header_row,
        group_test_header_row + len(comparison.group_effect_tests) + 1,
    ):
        horizontal = "center" if row_number == group_test_header_row else None
        summary.cell(row_number, 6).alignment = Alignment(
            horizontal=horizontal, vertical="center", wrap_text=True
        )
        summary.cell(row_number, 14).alignment = Alignment(
            horizontal=horizontal, vertical="center", wrap_text=True
        )
        summary.row_dimensions[row_number].height = 34

    details = workbook.create_sheet("配對明細")
    details.sheet_view.showGridLines = False
    detail_headers = [
        "受試者",
        "前測受試者",
        "後測受試者",
        "群組",
        "前測來源報表",
        "後測來源報表",
        f"前測{comparison.metric.display_name}",
        f"後測{comparison.metric.display_name}",
        "後測減前測",
        "成對資料",
        "自動排除原因",
    ]
    details.append(detail_headers)
    for row in comparison.rows:
        difference = (
            row.second_value - row.first_value
            if row.first_value is not None and row.second_value is not None
            else None
        )
        details.append(
            [
                row.display_name,
                row.first_subject,
                row.second_subject,
                row.group,
                row.first_report,
                row.second_report,
                row.first_value,
                row.second_value,
                difference,
                "是" if difference is not None else "否",
                row.exclusion_reason,
            ]
        )
    _style_header(details, 1, 1, len(detail_headers))
    details.freeze_panes = "A2"
    details.auto_filter.ref = details.dimensions
    for row in details.iter_rows(min_row=2, min_col=7, max_col=9):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"
    _fit_columns(details)

    trends = workbook.create_sheet("分段趨勢")
    trends.sheet_view.showGridLines = False
    trend_headers = [
        "受試者",
        "前測受試者",
        "後測受試者",
        "群組",
        "前測來源報表",
        "後測來源報表",
        "段落",
        "前測值",
        "後測值",
        "後測減前測",
        "自動排除原因",
    ]
    trends.append(trend_headers)
    segment_pool: dict[int, tuple[list[float], list[float]]] = {}
    for row in comparison.rows:
        first_segments = dict(row.first_segment_values)
        second_segments = dict(row.second_segment_values)
        for segment_index in sorted(set(first_segments) | set(second_segments)):
            first_value = first_segments.get(segment_index)
            second_value = second_segments.get(segment_index)
            difference = (
                second_value - first_value
                if first_value is not None and second_value is not None
                else None
            )
            trends.append(
                [
                    row.display_name,
                    row.first_subject,
                    row.second_subject,
                    row.group,
                    row.first_report,
                    row.second_report,
                    segment_index,
                    first_value,
                    second_value,
                    difference,
                    row.exclusion_reason,
                ]
            )
            if row.group != "排除":
                first_values, second_values = segment_pool.setdefault(segment_index, ([], []))
                if first_value is not None:
                    first_values.append(first_value)
                if second_value is not None:
                    second_values.append(second_value)
    _style_header(trends, 1, 1, len(trend_headers))
    trends.freeze_panes = "A2"
    trends.auto_filter.ref = f"A1:K{max(trends.max_row, 1)}"
    for row in trends.iter_rows(min_row=2, min_col=8, max_col=10):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.000000"

    average_header_row = 1
    trends.cell(average_header_row, 12, "段落")
    trends.cell(average_header_row, 13, comparison.first_source)
    trends.cell(average_header_row, 14, comparison.second_source)
    for offset, segment_index in enumerate(sorted(segment_pool), 1):
        first_values, second_values = segment_pool[segment_index]
        trends.cell(average_header_row + offset, 12, segment_index)
        trends.cell(
            average_header_row + offset,
            13,
            sum(first_values) / len(first_values) if first_values else None,
        )
        trends.cell(
            average_header_row + offset,
            14,
            sum(second_values) / len(second_values) if second_values else None,
        )
    _style_header(trends, average_header_row, 12, 14)
    if segment_pool:
        trend_chart = LineChart()
        trend_chart.style = 13
        trend_chart.title = f"各段{comparison.metric.display_name}平均趨勢"
        trend_chart.y_axis.title = comparison.metric.display_name
        trend_chart.x_axis.title = "段落"
        trend_chart.height = 8
        trend_chart.width = 14
        trend_chart.add_data(
            Reference(
                trends,
                min_col=13,
                max_col=14,
                min_row=average_header_row,
                max_row=average_header_row + len(segment_pool),
            ),
            titles_from_data=True,
        )
        trend_chart.set_categories(
            Reference(
                trends,
                min_col=12,
                min_row=average_header_row + 1,
                max_row=average_header_row + len(segment_pool),
            )
        )
        trends.add_chart(trend_chart, "L7")
    _fit_columns(trends)

    sources = workbook.create_sheet("來源清單")
    sources.sheet_view.showGridLines = False
    sources.append(["測驗階段", "來源 Excel"])
    for source in comparison.first_files:
        sources.append(["前測", source])
    for source in comparison.second_files:
        sources.append(["後測", source])
    _style_header(sources, 1, 1, 2)
    sources.freeze_panes = "A2"
    sources.auto_filter.ref = sources.dimensions
    _fit_columns(sources, maximum=80)

    info = workbook.create_sheet("報表資訊")
    info.sheet_view.showGridLines = False
    for row in [
        ("欄位", "值"),
        ("schema_version", SCHEMA_VERSION),
        ("report_type", "experiment_comparison"),
        ("generated_at", datetime.now().astimezone().isoformat(timespec="seconds")),
        ("distance_metric", comparison.metric.value),
        ("pairing_mode", comparison.pairing_mode.value),
        ("first_source", comparison.first_source),
        ("second_source", comparison.second_source),
        ("first_files", "\n".join(comparison.first_files)),
        ("second_files", "\n".join(comparison.second_files)),
    ]:
        info.append(row)
    _style_header(info, 1, 1, 2)
    _fit_columns(info, maximum=52)
    info.column_dimensions["A"].width = 26
    info.column_dimensions["B"].width = 52

    _save_atomic(workbook, path)
